"""Diagram association, content retention, and slide geometry regressions."""
import shutil
import sys
import unittest
import uuid
from pathlib import Path
from unittest.mock import patch

from PIL import Image
from pptx import Presentation

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'src'))
from pipeline.slides import scenes, export_slides
from pipeline.exporter import _body_groups


class SlideLayoutTest(unittest.TestCase):
    def test_diagrams_keep_numbered_prompt_and_explanation(self):
        groups = [('text', ['1. Previous item']), ('text', ['2. Composite functions', '2. 复合函数']),
                  ('text', ['Formula explanation']), ('fig', 'wide.png'), ('fig', 'second.png'),
                  ('text', ['3. Next item'])]
        result = list(scenes(groups))
        self.assertEqual(result[0], (['1. Previous item'], None))
        self.assertEqual(result[1][0], ['2. Composite functions\n2. 复合函数', 'Formula explanation'])
        self.assertEqual(result[1][0], result[2][0])
        self.assertEqual(result[-1], (['3. Next item'], None))

    def test_wide_and_portrait_figures_share_slides_with_editable_text(self):
        root = Path(__file__).resolve().parents[2] / 'dist' / 'build' / 'tests' / ('slides_' + uuid.uuid4().hex)
        root.mkdir(parents=True)
        self.addCleanup(shutil.rmtree, root)
        Image.new('RGB', (900, 280), 'white').save(root / 'wide.png')
        Image.new('RGB', (280, 600), 'white').save(root / 'portrait.png')
        blocks = []
        for idx, name in enumerate(('wide.png', 'portrait.png'), 1):
            blocks.extend([{'t': 'item', 'zh': f'{idx}. 题目与公式', 'en': f'{idx}. Question and formula', 'sub': {}},
                           {'t': 'fig', 'file': name}])
        doc = {'title': 'Fixture', 'pages': [{'number': 1, 'resolved': {'title_zh': '函数', 'title_en': 'Functions', 'blocks': blocks}}]}
        with patch('pipeline.slides.converter', return_value=None):
            export_slides(doc, 'bi', _body_groups, root, root / 'test.pptx', lambda: False)
        deck = Presentation(root / 'test.pptx')
        self.assertEqual(len(deck.slides), 2)
        for slide in deck.slides:
            diagram = next(s for s in slide.shapes if s.shape_type == 13)
            body = next(s for s in slide.shapes if s.name == 'Body')
            self.assertIn('题目与公式\n', body.text)
            self.assertIn('Question and formula', body.text)
            separated = (body.left + body.width <= diagram.left or diagram.left + diagram.width <= body.left
                         or body.top + body.height <= diagram.top or diagram.top + diagram.height <= body.top)
            self.assertTrue(separated)
            self.assertLessEqual(diagram.left + diagram.width, deck.slide_width)
            self.assertLessEqual(diagram.top + diagram.height, deck.slide_height)


if __name__ == '__main__':
    unittest.main()
