"""Full conversion gating: no partial output, optional LLM, automatic formats."""
import json
import shutil
import sys
import unittest
import uuid
from pathlib import Path
from unittest.mock import patch

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'src'))
from pipeline import convert, workspace


class ConvertTest(unittest.TestCase):
    def setUp(self):
        self.root = Path(__file__).resolve().parents[2]/'dist'/'build'/'tests'/('convert_'+uuid.uuid4().hex)
        self.root.mkdir(parents=True)
        self.addCleanup(shutil.rmtree, self.root)
        work = patch.object(workspace, 'WORK_ROOT', str(self.root/'work'))
        work.start();self.addCleanup(work.stop)
        self.source = self.root/'sample.pdf'
        self.source.write_bytes(b'%PDF-fixture-cached')
        self.destination = self.root/'output'
        self.doc = {'schema_version': 2, 'analysis_version': 3, 'key': 'sample', 'title': 'Sample',
                    'source': {'sha256': workspace.fingerprint(self.source), 'path': str(self.source)},
                    'pages': [{'number': 1, 'width': 400, 'height': 600, 'native_text': 'An English paragraph.',
                               'native_language': 'en', 'issues': [], 'resolved': None, 'figures': [], 'words': []}]}
        workspace.save_doc('sample', self.doc)
        self.logs = []

    def fake_build(self, key, target, formats, **kwargs):
        self.assertEqual(set(formats), {'tex', 'pdf', 'pptx'})
        stage = Path(kwargs['output_dir'])
        for ext in formats:
            (stage/f'{key}_{target.upper()}.{ext}').write_text('test artifact')

    def test_complete_local_content_never_calls_llm_and_outputs_all_formats(self):
        with patch.object(convert, 'complete_api') as api, patch.object(convert, 'build_v2', side_effect=self.fake_build):
            result = convert.convert_pdf(self.source, self.destination, 'en', True, progress=self.logs.append)
        api.assert_not_called()
        self.assertFalse(result['used_llm'])
        self.assertEqual({p.suffix for p in self.destination.iterdir()}, {'.tex', '.pdf', '.pptx'})
        self.assertTrue(any('不需要连接 LLM' in line for line in self.logs))

    def test_selected_formats_all_languages_and_pptx_without_tex_installation(self):
        self.set_bilingual()
        for selection, expected in [('pptx', ('pptx',)), ('tex-pdf', ('tex', 'pdf'))]:
            destination = self.root / selection
            def build(key, target, formats, **kwargs):
                self.assertEqual(formats, expected)
                for ext in formats:
                    (Path(kwargs['output_dir']) / f'{key}_{target.upper()}.{ext}').write_text('artifact')
            with patch.object(convert, 'build_v2', side_effect=build), patch.object(convert.shutil, 'which', return_value=None):
                result = convert.convert_pdf(self.source, destination, 'all', output_format=selection, progress=self.logs.append)
            self.assertEqual(len(result['files']), len(expected) * 3)
            self.assertEqual({p.suffix for p in destination.iterdir()}, {'.' + ext for ext in expected})

    def test_pptx_model_completion_does_not_require_xelatex(self):
        def finish(*args, **kwargs):
            self.set_bilingual()
            return 1
        def build(key, target, formats, **kwargs):
            self.assertEqual(formats, ('pptx',))
            (Path(kwargs['output_dir']) / f'{key}_{target.upper()}.pptx').write_text('artifact')
        with patch.object(convert, 'complete_api', side_effect=finish), patch.object(convert.shutil, 'which', return_value=None), patch.object(convert, 'build_v2', side_effect=build):
            result = convert.convert_pdf(self.source, self.destination, 'bi', True, output_format='pptx', progress=self.logs.append)
        self.assertTrue(result['used_llm'])

    def test_incomplete_without_checkbox_stops_before_any_output_or_network(self):
        with patch.object(convert, 'complete_api') as api, patch.object(convert, 'build_v2') as build:
            with self.assertRaises(workspace.NeedsModel):
                convert.convert_pdf(self.source, self.destination, 'bi', progress=self.logs.append)
        api.assert_not_called();build.assert_not_called()
        self.assertFalse(self.destination.exists())

    def test_checkbox_automatically_completes_then_outputs_without_second_click(self):
        def finish(*args, **kwargs):
            self.doc['pages'][0]['resolved'] = {'needs_review': False, 'blocks': [
                {'t': 'item', 'en': 'An English paragraph.', 'zh': '一段英文。', 'sub': {'en': [], 'zh': []}}]}
            workspace.save_doc('sample', self.doc)
            return 1
        with patch.object(convert.shutil, 'which', return_value='xelatex'), \
             patch.object(convert, 'complete_api', side_effect=finish) as api, \
             patch.object(convert, 'build_v2', side_effect=self.fake_build):
            result = convert.convert_pdf(self.source, self.destination, 'bi', True, 'https://example.com/v1', 'model', progress=self.logs.append)
        api.assert_called_once()
        self.assertTrue(result['used_llm'])
        self.assertEqual(len(result['files']), 3)

    def test_failed_build_does_not_publish_tex_or_replace_previous_output(self):
        self.destination.mkdir()
        previous = self.destination/'sample_EN.pdf'
        previous.write_bytes(b'previous-output')
        def failure(*args, **kwargs):
            (Path(kwargs['output_dir'])/'sample_EN.tex').write_text('partial')
            raise RuntimeError('compiler failed')
        with patch.object(convert, 'build_v2', side_effect=failure):
            with self.assertRaisesRegex(RuntimeError, 'compiler failed'):
                convert.convert_pdf(self.source, self.destination, 'en', progress=self.logs.append)
        self.assertEqual(previous.read_bytes(), b'previous-output')
        self.assertFalse((self.destination/'sample_EN.tex').exists())

    def test_llm_uncertainty_still_blocks_final_output(self):
        with patch.object(convert.shutil, 'which', return_value='xelatex'), \
             patch.object(convert, 'complete_api', return_value=1), patch.object(convert, 'build_v2') as build:
            with self.assertRaises(workspace.NeedsModel):
                convert.convert_pdf(self.source, self.destination, 'bi', True, 'https://example.com', 'model', progress=self.logs.append)
        build.assert_not_called()
        self.assertFalse(self.destination.exists())

    def set_bilingual(self):
        self.doc['pages'][0]['resolved'] = {'needs_review': False, 'blocks': [
            {'t': 'item', 'zh': '中文第一句。', 'en': 'First English sentence.',
             'sub': {'zh': ['中文子项。'], 'en': ['English subitem.']}}]}
        workspace.save_doc('sample', self.doc)

    def test_all_builds_nine_files_without_any_model_for_complete_bilingual(self):
        self.set_bilingual()
        with patch.object(convert, 'complete_api') as api, patch.object(convert, 'build_v2', side_effect=self.fake_build) as build:
            result = convert.convert_pdf(self.source, self.destination, 'all', progress=self.logs.append)
        api.assert_not_called()
        self.assertEqual([call.args[1] for call in build.call_args_list], ['zh', 'en', 'bi'])
        self.assertEqual(len(result['files']), 9)
        self.assertEqual(len(list(self.destination.glob('sample_*.*'))), 9)

    def test_all_uses_one_bilingual_model_pass(self):
        def finish(*args, **kwargs):
            self.assertEqual(args[1], 'bi')
            self.set_bilingual()
            return 1
        with patch.object(convert.shutil, 'which', return_value='xelatex'), \
             patch.object(convert, 'complete_api', side_effect=finish) as api, \
             patch.object(convert, 'build_v2', side_effect=self.fake_build):
            result = convert.convert_pdf(self.source, self.destination, 'all', True, 'https://example.com', 'model', progress=self.logs.append)
        api.assert_called_once()
        self.assertEqual(len(result['files']), 9)

    def test_all_does_not_publish_partial_languages_when_last_build_fails(self):
        self.set_bilingual()
        def build(key, target, formats, **kwargs):
            if target == 'bi':
                raise RuntimeError('bilingual compile failed')
            self.fake_build(key, target, formats, **kwargs)
        with patch.object(convert, 'build_v2', side_effect=build):
            with self.assertRaisesRegex(RuntimeError, 'bilingual compile failed'):
                convert.convert_pdf(self.source, self.destination, 'all', progress=self.logs.append)
        self.assertFalse(self.destination.exists())

    def test_bilingual_chinese_then_english_for_body_and_subitems(self):
        self.set_bilingual()
        from pipeline.exporter import _body_lines, build_v2
        lines = [text for kind, text in _body_lines(self.doc['pages'][0], 'bi')]
        self.assertEqual(lines, ['中文第一句。', 'First English sentence.', '中文子项。', 'English subitem.'])
        stage = self.root/'bilingual'
        build_v2('sample', 'bi', ('tex', 'pptx'), progress=lambda _: None, output_dir=stage)
        tex = (stage/'sample_BI.tex').read_text(encoding='utf-8')
        self.assertLess(tex.index('中文第一句。'), tex.index('First English sentence.'))
        self.assertLess(tex.index('中文子项。'), tex.index('English subitem.'))
        self.assertIn(r'\nopagebreak[4]', tex)
        from pptx import Presentation
        deck = Presentation(stage/'sample_BI.pptx')
        body = '\n'.join(shape.text for slide in deck.slides for shape in slide.shapes if shape.has_text_frame)
        self.assertIn('中文第一句。\nFirst English sentence.', body)
        self.assertIn('中文子项。\nEnglish subitem.', body)


if __name__ == '__main__':
    unittest.main()
