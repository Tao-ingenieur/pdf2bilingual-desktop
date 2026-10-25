"""Compose editable text and its source diagrams on the same slide."""
import math
import re
import unicodedata

from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_AUTO_SIZE

from .office_math import converter, set_rich_text
from .tex2plain import tex2plain


def text_height(text, width, size=20):
    """Conservative line estimate in inches, including native equation height."""
    capacity = max(8, (width * 72 - 15) / size)
    height = 0
    for line in text.split('\n'):
        plain = tex2plain(line)
        units = sum(1.05 if unicodedata.east_asian_width(c) in ('W', 'F') else .56 for c in plain)
        rows = max(1, math.ceil(units / capacity))
        complex_math = bool(re.search(r'\\(?:frac|dfrac|binom|begin|sqrt)', line))
        height += rows * size / 72 * (2.1 if complex_math else 1.7)
    return height + .12


def scenes(groups):
    """Attach each figure to the preceding numbered item and its explanation.

    Unnumbered data uses the immediately preceding paragraph. Leading figures
    use the following paragraph. Consecutive diagrams retain the same context.
    """
    groups = list(groups)
    buffered = []
    context = []
    for index, (kind, value) in enumerate(groups):
        if kind != 'fig':
            buffered.append('\n'.join(value))
            continue
        if buffered:
            start = len(buffered) - 1
            for pos in range(len(buffered) - 1, -1, -1):
                if re.match(r'^\s*\d+[.、．\s]', buffered[pos]):
                    start = pos
                    break
            if start:
                yield buffered[:start], None
            context = buffered[start:]
            buffered = []
        if not context:
            context = next((['\n'.join(v)] for k, v in groups[index+1:] if k != 'fig'), [])
        yield list(context), value
    if buffered:
        yield buffered, None


def paginate(texts, width, height):
    """Split between aligned bilingual groups, never through a formula."""
    current, used = [], 0
    for text in texts:
        needed = text_height(text, width)
        if current and used + needed > height - .15:
            yield '\n\n'.join(current)
            current, used = [], 0
        current.append(text)
        # An empty paragraph between groups has its own line height in Office.
        used += needed + .5
    if current:
        yield '\n\n'.join(current)
    elif not texts:
        yield ''


def export_slides(doc, lang, groups_for_page, figures, destination, cancelled):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    native_math = converter() is not None

    def textbox(slide, text, x, y, w, h, size, color='252F33', bold=False, math_text=False):
        shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        shape.name = 'Body' if math_text else 'Heading' if bold else 'Source'
        frame = shape.text_frame
        frame.word_wrap = True
        frame.margin_left = frame.margin_right = Inches(.03)
        frame.margin_top = frame.margin_bottom = Inches(.02)
        frame.text = text
        if math_text and native_math:
            set_rich_text(frame, text, size=size)
        frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
        for paragraph in frame.paragraphs:
            paragraph.font.name = 'Microsoft YaHei'
            paragraph.font.size = Pt(size)
            paragraph.font.bold = bold
            paragraph.font.color.rgb = RGBColor.from_string(color)
            paragraph.space_after = Pt(4)
        return shape

    def picture(slide, filename, x, y, w, h):
        with Image.open(figures / filename) as image:
            ratio = image.width / image.height
        width = min(w, h * ratio)
        height = width / ratio
        shape = slide.shapes.add_picture(str(figures / filename), Inches(x + (w-width)/2),
                                        Inches(y + (h-height)/2), width=Inches(width), height=Inches(height))
        shape.name = 'Related diagram'

    for page in doc['pages']:
        if cancelled():
            raise InterruptedError('生成已取消。')
        if lang == 'source':
            slide = prs.slides.add_slide(prs.slide_layouts[6])
            picture(slide, page['preview'], 0, 0, 13.333, 7.5)
            continue
        titles = page.get('resolved') or {}
        heading = ('\n'.join(titles.get('title_' + l, '') for l in ('zh', 'en')) if lang == 'bi'
                   else titles.get('title_' + lang, ''))
        heading = heading.strip() or f'{doc["title"]} · {page["number"]}'
        for texts, filename in scenes(groups_for_page(page, lang)):
            # Landscape diagrams follow the text, portrait/square diagrams sit beside it.
            wide = False
            if filename:
                with Image.open(figures / filename) as image:
                    wide = image.width / image.height >= 1.7
            width, height = ((12.1, 2.5) if wide else (6.0, 5.25)) if filename else (12.1, 5.25)
            for text in paginate(texts, width, height):
                if cancelled():
                    raise InterruptedError('生成已取消。')
                slide = prs.slides.add_slide(prs.slide_layouts[6])
                title_units = max(sum(1 if unicodedata.east_asian_width(c) in ('W', 'F') else .56
                                      for c in line) for line in heading.split('\n'))
                title_size = 23 if lang == 'bi' else 24 if title_units > 26 else 30
                textbox(slide, heading, .45, .18, 12.4, .95, title_size, '117D76', True)
                text = text if native_math else tex2plain(text)
                if wide:
                    # Give a short prompt more room for its diagram, as in the reference.
                    text_h = min(2.5, max(.7, text_height(text, width)))
                    textbox(slide, text, .6, 1.35, width, text_h, 20, math_text=True)
                    picture(slide, filename, .65, 1.55 + text_h, 12.0, 5.25 - text_h)
                else:
                    textbox(slide, text, .6, 1.35, width, height, 20, math_text=True)
                    if filename:
                        picture(slide, filename, 7.0, 1.4, 5.7, 5.15)
                source = f'原件第 {page["number"]} 页' if lang in ('zh', 'bi') else f'Source page {page["number"]}'
                textbox(slide, source, .6, 7.02, 12.1, .28, 11, '61717B')
    prs.save(str(destination))
    return len(prs.slides)
