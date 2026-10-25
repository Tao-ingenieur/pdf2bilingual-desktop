"""Optional native PowerPoint equations using the local Office MathML transform."""
import os
import re
from functools import lru_cache
from pathlib import Path

M = 'http://schemas.openxmlformats.org/officeDocument/2006/math'
A = 'http://schemas.openxmlformats.org/drawingml/2006/main'
A14 = 'http://schemas.microsoft.com/office/drawing/2010/main'
W = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'


@lru_cache(maxsize=1)
def converter():
    try:
        from latex2mathml.converter import convert
        from lxml import etree
    except ImportError:
        return None
    override = os.environ.get('PDFBILINGUAL_MML2OMML')
    paths = [Path(override)] if override else [
        Path(os.environ.get(root, 'C:/Program Files')) / 'Microsoft Office/root/Office16/MML2OMML.XSL'
        for root in ('ProgramFiles', 'ProgramFiles(x86)')]
    path = next((p for p in paths if p.is_file()), None)
    if path is None:
        return None
    return convert, etree.XSLT(etree.parse(str(path)))


@lru_cache(maxsize=4096)
def equation(source):
    from lxml import etree
    convert, transform = converter()
    # PowerPoint does not reliably stretch delimiters around a one-column
    # matrix imported through MathML. Use the equivalent transposed row vector.
    def column_vector(match):
        body = match.group(1)
        if '&' in body:
            return match.group(0)
        components = re.split(r'\\\\', body)
        return r'\left(' + ','.join(components) + r'\right)^{\mathrm{T}}'
    source = re.sub(r'\\begin\{pmatrix\}(.*?)\\end\{pmatrix\}', column_vector, source, flags=re.S)
    mathml = etree.fromstring(convert(source).encode('utf-8'))
    root = transform(mathml).getroot()
    if root.tag != f'{{{M}}}oMath':
        root = root.find(f'.//{{{M}}}oMath')
    if root is None:
        raise ValueError('公式转换未产生 Office Math。')
    for props in root.iter(f'{{{M}}}dPr'):
        if props.find(f'{{{M}}}begChr') is None:
            props.insert(0, etree.Element(f'{{{M}}}begChr', {f'{{{M}}}val': '('}))
        if props.find(f'{{{M}}}endChr') is None:
            props.insert(len(props), etree.Element(f'{{{M}}}endChr', {f'{{{M}}}val': ')'}))
        etree.SubElement(props, f'{{{M}}}grow', {f'{{{M}}}val': '1'})
        etree.SubElement(props, f'{{{M}}}shp', {f'{{{M}}}val': 'centered'})
        control = etree.SubElement(props, f'{{{M}}}ctrlPr')
        rp = etree.SubElement(control, f'{{{A}}}rPr', sz='1800')
        etree.SubElement(rp, f'{{{A}}}latin', typeface='Cambria Math')
    # DrawingML equations use DrawingML run properties, not Word run properties.
    for node in list(root.iter(f'{{{W}}}rPr')):
        node.getparent().remove(node)
    for run in root.iter(f'{{{M}}}r'):
        props = etree.Element(f'{{{A}}}rPr', sz='1800', lang='en-US')
        etree.SubElement(props, f'{{{A}}}latin', typeface='Cambria Math')
        run.insert(0, props)
    return etree.tostring(root)


def set_rich_text(frame, text, size=18):
    """Return False if the optional native-equation backend is unavailable."""
    if converter() is None:
        return False
    from lxml import etree
    from pptx.util import Pt
    frame.clear()
    for i, line in enumerate(text.splitlines()):
        paragraph = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        paragraph.font.name = 'Microsoft YaHei'
        paragraph.font.size = Pt(size)
        for part in re.split(r'(\$\$.*?\$\$|\$[^$]*\$)', line):
            if part.startswith('$') and part.endswith('$'):
                node = etree.Element(f'{{{A14}}}m', nsmap={'a14': A14, 'm': M})
                node.append(etree.fromstring(equation(part.strip('$'))))
                paragraph._p.append(node)
            elif part:
                run = paragraph.add_run()
                run.text = part
    return True
