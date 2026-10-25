# converter/pptxgen.py — 生成 PPTX
# ---------------- PPT ----------------
import os
from pdf2bilingual.paths import NAVY, LIGHT, GREY, WHITE, EN_FONT, ZH_FONT, WORK_ROOT
from .tex2plain import tex2plain
from pptx.util import Inches, Pt, Emu
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor


def ppt_color(rgb):
    return RGBColor(*rgb)


def add_rect(slide, l, t, w, h, fill_rgb, line_rgb=None):
    sp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, l, t, w, h)
    sp.fill.solid(); sp.fill.fore_color.rgb = ppt_color(fill_rgb)
    if line_rgb is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = ppt_color(line_rgb); sp.line.width = 9525
    sp.shadow.inherit = False
    return sp


def add_text(slide, l, t, w, h, text, size=18, bold=False, color=GREY,
             font=EN_FONT, align="left", anchor="top"):
    tb = slide.shapes.add_textbox(l, t, w, h)
    tf = tb.text_frame; tf.word_wrap = True
    tf.vertical_anchor = (MSO_ANCHOR.MIDDLE if anchor == "middle"
                          else MSO_ANCHOR.TOP)
    tf.margin_left = 91440; tf.margin_right = 91440
    lines = text.split("\n")
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = (PP_ALIGN.CENTER if align == "center"
                       else PP_ALIGN.LEFT)
        r = p.add_run(); r.text = ln
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.color.rgb = ppt_color(color)
        r.font.name = font
    return tb


def blank_slide(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def ppt_cover(prs, doc, lang):
    s = blank_slide(prs)
    add_rect(s, 0, 0, prs.slide_width, prs.slide_height, NAVY)
    if lang == "bi":
        add_text(s, Inches(0.8), Inches(2.4), Inches(11.7), Inches(1.0),
                 doc["title_en"], size=38, bold=True, color=WHITE,
                 font=EN_FONT, align="center")
        add_text(s, Inches(0.8), Inches(3.5), Inches(11.7), Inches(0.8),
                 doc["title_zh"], size=24, bold=True, color=(210, 220, 235),
                 font=ZH_FONT, align="center")
        add_text(s, Inches(0.8), Inches(4.5), Inches(11.7), Inches(0.6),
                 f"{doc['subtitle_en']}  ·  {doc['subtitle_zh']}", size=16,
                 color=(200, 210, 230), font=ZH_FONT, align="center")
        series = "Bilingual Edition · 双语版"
    else:
        title = doc[f"title_{lang}"]
        sub = doc[f"subtitle_{lang}"]
        font = ZH_FONT if lang == "zh" else EN_FONT
        add_text(s, Inches(0.8), Inches(2.7), Inches(11.7), Inches(1.4),
                 title, size=40, bold=True, color=WHITE, font=font, align="center")
        add_text(s, Inches(0.8), Inches(4.1), Inches(11.7), Inches(0.8),
                 sub, size=20, color=(210, 220, 235), font=font, align="center")
        series = doc.get(f"series_{lang}") or ("Bilingual Edition" if lang == "en" else "双语版")
    add_text(s, Inches(0.8), Inches(6.6), Inches(11.7), Inches(0.5),
             series, size=13, color=(180, 195, 215), font=ZH_FONT, align="center")
    return s


def ppt_divider(prs, doc, part, lang):
    s = blank_slide(prs)
    add_rect(s, 0, 0, prs.slide_width, prs.slide_height, NAVY)
    add_rect(s, Inches(4.4), Inches(4.5), Inches(4.5), Inches(0.06), (200, 215, 235))
    p = [x for x in doc["parts"] if x["id"] == part][0]
    if lang == "bi":
        add_text(s, Inches(1), Inches(2.9), Inches(11.3), Inches(1.0),
                 p["title_en"], size=34, bold=True, color=WHITE,
                 font=EN_FONT, align="center")
        add_text(s, Inches(1), Inches(4.0), Inches(11.3), Inches(0.9),
                 p["title_zh"], size=26, bold=True, color=(210, 220, 235),
                 font=ZH_FONT, align="center")
    else:
        add_text(s, Inches(1), Inches(3.1), Inches(11.3), Inches(1.2),
                 p[f"title_{lang}"], size=36, bold=True, color=WHITE,
                 font=(ZH_FONT if lang == "zh" else EN_FONT), align="center")
    return s


def _place_fig(slide, fb, x, y, w, h, key=None, lang=None):
    from PIL import Image
    figdir = os.path.join(WORK_ROOT, key, "figures")
    path = os.path.join(figdir, fb["file"])
    if not os.path.exists(path):
        return
    try:
        iw, ih = Image.open(path).size
    except Exception:
        iw, ih = 1000, 1000
    ar = iw / ih; bar = w / h
    if ar > bar:
        nw = w; nh = int(w / ar)
    else:
        nh = h; nw = int(h * ar)
    nx = x + (w - nw) // 2; ny = y + (h - nh) // 2
    slide.shapes.add_picture(path, nx, ny, nw, nh)


def ppt_table(slide, rows, y, font):
    from pptx.util import Inches, Pt
    nrows = len(rows); ncols = len(rows[0])
    gtbl = slide.shapes.add_table(nrows, ncols, Inches(0.5), y,
                                  Inches(12.3), Inches(2.3))
    table = gtbl.table
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            c = table.cell(ri, ci)
            c.text = val
            for para in c.text_frame.paragraphs:
                for run in para.runs:
                    run.font.size = Pt(11); run.font.name = font
                    run.font.color.rgb = ppt_color(GREY)
            c.fill.solid(); c.fill.fore_color.rgb = ppt_color(WHITE)
            c.margin_left = 91440; c.margin_right = 91440
            c.margin_top = 30000; c.margin_bottom = 30000
    for ci in range(ncols):
        c = table.cell(0, ci)
        c.fill.solid(); c.fill.fore_color.rgb = ppt_color((217, 226, 239))
        for para in c.text_frame.paragraphs:
            for run in para.runs:
                run.font.bold = True; run.font.color.rgb = ppt_color(NAVY)
    ri = 1
    while ri < nrows:
        if rows[ri][0].strip() == "":
            j = ri - 1
            while j >= 0 and rows[j][0].strip() == "":
                j -= 1
            if j >= 0:
                table.cell(j, 0).merge(table.cell(ri, 0))
        ri += 1


def ppt_chapter(prs, doc, ch, idx, lang, key=None):
    s = blank_slide(prs)
    add_rect(s, 0, 0, prs.slide_width, Inches(1.0), NAVY)
    add_rect(s, Inches(0.35), Inches(0.22), Inches(0.6), Inches(0.6), WHITE)
    add_text(s, Inches(0.35), Inches(0.22), Inches(0.6), Inches(0.6),
             str(idx), size=22, bold=True, color=NAVY, align="center", anchor="middle")
    if lang == "bi":
        title = f"{ch['title_en']}  ·  {ch['title_zh']}"
        font = ZH_FONT
    else:
        title = ch["title_en"] if lang == "en" else ch["title_zh"]
        font = ZH_FONT if lang == "zh" else EN_FONT
    add_text(s, Inches(1.1), Inches(0.18), Inches(11.4), Inches(0.66),
             title, size=20, bold=True, color=WHITE, font=font, anchor="middle")

    body = ch["blocks"]
    is_proofs = bool(body and body[0].get("t") == "proofs")
    has_fig = any(b["t"] == "fig" for b in body)
    has_table = any(b["t"] == "table" for b in body)
    top = Inches(1.25)
    if is_proofs:
        text_avail = Inches(5.6)
    elif has_table:
        text_avail = Inches(3.0)
    elif has_fig:
        text_avail = Inches(3.4)
    else:
        text_avail = Inches(5.6)

    add_rect(s, Inches(0.5), top, Inches(12.3), text_avail, LIGHT)
    tf_box = add_text(s, Inches(0.8), top + Inches(0.1), Inches(11.7),
                      text_avail - Inches(0.2), "", size=14)
    tf = tf_box.text_frame; tf.word_wrap = True

    if is_proofs:
        if lang == "bi":
            intro_en = ch.get("intro_en", []); intro_zh = ch.get("intro_zh", [])
            if intro_en:
                p0 = tf.paragraphs[0]
                r = p0.add_run(); r.text = tex2plain("  ".join(intro_en))
                r.font.size = Pt(13); r.font.color.rgb = ppt_color(GREY); r.font.name = font
            if intro_zh:
                pz = tf.add_paragraph()
                rz = pz.add_run(); rz.text = tex2plain("  ".join(intro_zh))
                rz.font.size = Pt(12); rz.font.color.rgb = ppt_color((130, 130, 130))
                rz.font.name = ZH_FONT
            for pr in body[0]["proofs"]:
                pp = tf.add_paragraph()
                r = pp.add_run()
                r.text = f"● {tex2plain(pr['name_en'])}   {tex2plain(pr['formula_en'])}"
                r.font.size = Pt(13); r.font.bold = True
                r.font.color.rgb = ppt_color(NAVY); r.font.name = font
                for st in pr["steps_en"]:
                    p2 = tf.add_paragraph()
                    r2 = p2.add_run(); r2.text = f"    {tex2plain(st)}"
                    r2.font.size = Pt(12); r2.font.color.rgb = ppt_color((90, 90, 90))
                    r2.font.name = font
                ppz = tf.add_paragraph()
                rz = ppz.add_run()
                rz.text = f"● {tex2plain(pr['name_zh'])}   {tex2plain(pr['formula_zh'])}"
                rz.font.size = Pt(13); rz.font.bold = True
                rz.font.color.rgb = ppt_color(NAVY); rz.font.name = ZH_FONT
                for st in pr["steps_zh"]:
                    p2 = tf.add_paragraph()
                    r2 = p2.add_run(); r2.text = f"    {tex2plain(st)}"
                    r2.font.size = Pt(12); r2.font.color.rgb = ppt_color((140, 140, 140))
                    r2.font.name = ZH_FONT
        else:
            intro = ch.get(f"intro_{lang}", [])
            if intro:
                p0 = tf.paragraphs[0]
                r = p0.add_run(); r.text = tex2plain("  ".join(intro))
                r.font.size = Pt(13); r.font.color.rgb = ppt_color(GREY); r.font.name = font
            for pr in body[0]["proofs"]:
                name = tex2plain(pr[f"name_{lang}"])
                formula = tex2plain(pr[f"formula_{lang}"])
                steps = [tex2plain(st) for st in pr[f"steps_{lang}"]]
                pp = tf.add_paragraph()
                r = pp.add_run(); r.text = f"● {name}   {formula}"
                r.font.size = Pt(13); r.font.bold = True
                r.font.color.rgb = ppt_color(NAVY); r.font.name = font
                for st in steps:
                    p2 = tf.add_paragraph()
                    r2 = p2.add_run(); r2.text = f"    {st}"
                    r2.font.size = Pt(12); r2.font.color.rgb = ppt_color((90, 90, 90))
                    r2.font.name = font
    else:
        num = 0
        for b in body:
            if b["t"] == "item":
                num += 1
                if lang == "bi":
                    en_txt = tex2plain(b["en"]); zh_txt = tex2plain(b["zh"])
                    en_sub = [tex2plain(x) for x in b.get("sub", {}).get("en", [])]
                    zh_sub = [tex2plain(x) for x in b.get("sub", {}).get("zh", [])]
                    first = (num == 1 and not tf.paragraphs[0].runs)
                    p = tf.paragraphs[0] if first else tf.add_paragraph()
                    r = p.add_run(); r.text = f"{num}. {en_txt}"
                    r.font.size = Pt(14); r.font.color.rgb = ppt_color(GREY); r.font.name = font
                    pz = tf.add_paragraph()
                    rz = pz.add_run(); rz.text = f"    {zh_txt}"
                    rz.font.size = Pt(12); rz.font.color.rgb = ppt_color((130, 130, 130))
                    rz.font.name = ZH_FONT
                    for e, z in zip(en_sub, zh_sub):
                        p2 = tf.add_paragraph()
                        r2 = p2.add_run(); r2.text = f"    • {e}"
                        r2.font.size = Pt(12); r2.font.color.rgb = ppt_color((90, 90, 90))
                        r2.font.name = font
                        pz2 = tf.add_paragraph()
                        rz2 = pz2.add_run(); rz2.text = f"      {z}"
                        rz2.font.size = Pt(11); rz2.font.color.rgb = ppt_color((150, 150, 150))
                        rz2.font.name = ZH_FONT
                else:
                    txt = tex2plain(b["en"] if lang == "en" else b["zh"])
                    sub = [tex2plain(x) for x in
                           (b["sub"]["en"] if lang == "en" else b["sub"]["zh"])]
                    first = (num == 1 and not tf.paragraphs[0].runs)
                    p = tf.paragraphs[0] if first else tf.add_paragraph()
                    r = p.add_run(); r.text = f"{num}. {txt}"
                    r.font.size = Pt(14); r.font.color.rgb = ppt_color(GREY); r.font.name = font
                    for it in sub:
                        p2 = tf.add_paragraph()
                        r2 = p2.add_run(); r2.text = f"    • {it}"
                        r2.font.size = Pt(12); r2.font.color.rgb = ppt_color((90, 90, 90))
                        r2.font.name = font
        if has_table:
            tbl_rows = [[tex2plain(c) for c in row] for row in
                        next(b["rows"] for b in body if b["t"] == "table")]
            ppt_table(s, tbl_rows, top + text_avail + Inches(0.1), font)

    if has_fig:
        figs = [b for b in body if b["t"] == "fig"]
        fig_y = top + text_avail + Inches(0.1)
        n = len(figs)
        if n == 1:
            _place_fig(s, figs[0], Inches(3.9), fig_y, Inches(5.0), Inches(2.6), key)
        elif n == 2:
            _place_fig(s, figs[0], Inches(0.6), fig_y, Inches(5.2), Inches(2.5), key)
            _place_fig(s, figs[1], Inches(6.7), fig_y, Inches(5.2), Inches(2.5), key)
        else:
            cw = Inches(3.9); chh = Inches(2.4)
            for i, fb in enumerate(figs):
                rr = i // 3; cc = i % 3
                _place_fig(s, fb, Inches(0.6) + cc * Inches(4.1),
                           fig_y + rr * Inches(2.6), cw, chh, key)

    add_text(s, Inches(0.5), Inches(7.05), Inches(8), Inches(0.4),
             (f"{doc['title_en']}  ·  {doc['title_zh']}" if lang == "bi"
              else doc[f"title_{lang}"]), size=9,
             color=(150, 150, 150), font=font)
    add_text(s, Inches(11.5), Inches(7.05), Inches(1.3), Inches(0.4),
             f"{idx}/{len(doc['chapters'])}", size=9, color=(150, 150, 150),
             font=font, align="right")
    return s


def build_ppt(doc, lang, key=None):
    from pptx import Presentation
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    _build_ppt_lang(prs, doc, lang, key)
    return prs


def _build_ppt_lang(prs, doc, lang, key):
    ppt_cover(prs, doc, lang)
    prev = None
    idx = 0
    for ch in doc["chapters"]:
        if ch["part"] != prev and prev is not None:
            ppt_divider(prs, doc, ch["part"], lang)
        prev = ch["part"]
        idx += 1
        ppt_chapter(prs, doc, ch, idx, lang, key)
    return prs